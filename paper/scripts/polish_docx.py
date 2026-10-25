"""Post-process the pandoc DOCX: bordered, compact tables and centred figures/table titles."""
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

d = docx.Document(sys.argv[1])
for t in d.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "808080")
        borders.append(e)
    tblPr.append(borders)
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "E8ECF7")
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.alignment = None
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.name = "Times New Roman"
                    if i == 0:
                        r.font.bold = True
for p in d.paragraphs:
    if p._p.xpath(".//w:drawing"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if p.text.startswith("Table ") and p.runs and p.runs[0].bold:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(9.5)
d.save(sys.argv[1])
print("polished", sys.argv[1])
